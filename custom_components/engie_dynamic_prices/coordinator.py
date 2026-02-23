"""Data coordinator for Engie Dynamic Prices."""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Optional

from homeassistant.core import HomeAssistant
from homeassistant.helpers.update_coordinator import DataUpdateCoordinator, UpdateFailed

from .const import DOMAIN, XLSX_URL

_LOGGER = logging.getLogger(__name__)

SCAN_INTERVAL = timedelta(hours=1)


@dataclass
class EngieData:
    """Parsed daily price data from Engie."""

    price_date: date
    prices: list[float]  # 24 values, index = hour (0-23), unit = EUR/kWh

    @property
    def current_price(self) -> float | None:
        hour = datetime.now().hour
        return self.prices[hour] if len(self.prices) > hour else None

    @property
    def next_price(self) -> float | None:
        next_hour = (datetime.now().hour + 1) % 24
        return self.prices[next_hour] if len(self.prices) > next_hour else None

    @property
    def average_price(self) -> float | None:
        return round(sum(self.prices) / len(self.prices), 6) if self.prices else None

    @property
    def min_price(self) -> float | None:
        return min(self.prices) if self.prices else None

    @property
    def max_price(self) -> float | None:
        return max(self.prices) if self.prices else None

    @property
    def cheapest_hours(self) -> list[dict]:
        """Return all hours sorted by ascending price."""
        return [
            {
                "hour": f"{i:02d}:00 - {(i + 1) % 24:02d}:00",
                "price_eur_kwh": p,
            }
            for i, p in sorted(enumerate(self.prices), key=lambda x: x[1])
        ]


class EngieCoordinator(DataUpdateCoordinator[EngieData]):
    """Coordinator that fetches and parses Engie dynamic energy prices."""

    def __init__(self, hass: HomeAssistant, session) -> None:
        super().__init__(
            hass,
            _LOGGER,
            name=DOMAIN,
            update_interval=SCAN_INTERVAL,
        )
        self._session = session
        self._last_fetch_date: Optional[date] = None

    async def _async_update_data(self) -> EngieData:
        # Return cached data if we already fetched today — no API call needed.
        # We key on *fetch date* (not price_date) so that day-ahead prices
        # published for tomorrow do not cause hourly re-fetches or prematurely
        # replace today's average / min / max with tomorrow's values.
        # The coordinator still runs hourly so that current_price / next_price
        # sensors reflect the correct hour.
        if self._last_fetch_date == date.today() and self.data is not None:
            _LOGGER.debug("Prices already fetched today (%s), skipping", date.today())
            return self.data

        try:
            async with self._session.get(XLSX_URL) as response:
                response.raise_for_status()
                content = await response.read()
        except Exception as err:
            raise UpdateFailed(f"Error fetching Engie prices: {err}") from err

        try:
            result = await self.hass.async_add_executor_job(
                EngieCoordinator._parse_xlsx, content
            )
        except Exception as err:
            raise UpdateFailed(f"Error parsing Engie prices: {err}") from err

        self._last_fetch_date = date.today()
        return result

    @staticmethod
    def _parse_xlsx(content: bytes) -> EngieData:
        import openpyxl  # imported here to avoid load-time cost

        wb = openpyxl.load_workbook(io.BytesIO(content))

        # Prefer the hourly sheet (name contains "Heures" or "Uren")
        hourly_sheet = next(
            (
                ws
                for ws in wb.worksheets
                if "Heures" in ws.title or "Uren" in ws.title
            ),
            wb.worksheets[0],
        )

        # Row 1, column 2 holds the price date as a datetime value
        date_cell = hourly_sheet.cell(row=1, column=2).value
        price_date = (
            date_cell.date() if isinstance(date_cell, datetime) else date.today()
        )

        # Data starts at row 5; column 2 holds the EUR/MWh price.
        # Stop at 24 values so that day-ahead rows published in the same
        # sheet do not skew average / min / max calculations.
        prices: list[float] = []
        for row in hourly_sheet.iter_rows(min_row=5, values_only=True):
            if len(prices) >= 24:
                break
            value = row[1]
            if isinstance(value, (int, float)):
                prices.append(round(float(value) / 1000, 6))  # EUR/MWh → EUR/kWh

        _LOGGER.debug(
            "Parsed %d hourly prices for %s", len(prices), price_date
        )
        return EngieData(price_date=price_date, prices=prices)
